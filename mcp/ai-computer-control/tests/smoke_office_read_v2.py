"""办公/文件「读」路径二期冒烟 —— 用真文件(临时目录里现造 docx/xlsx/pdf/pptx/GBK/UTF-16)逐项钉住:

  ① edit_file      —— 保留 CRLF 行尾,不双写 UTF-8 BOM,保留原 BOM;.bat/.ps1 仍是 CRLF(+BOM)
  ② read_file      —— 剥 UTF-8 BOM、解 UTF-16(BOM/无 BOM)、GBK 回退、二进制拒读;普通中文不再附 non_ascii 噪声
  ③ list_directory —— include_hidden=False 剪掉 .git;坏符号链接跳过并计数;total/capped 在 entries 之前;非递归也封顶
  ④ memory         —— 瞬时 OSError 不把健康库改名成 .corrupt;真解析失败才隔离
  ⑤ read_document  —— docx 保结构(标题/列表/表格顺序)+ 字符上限/续读;加密 PDF -> ok:false;扫描 PDF -> note;
                      加密 xlsx / 坏 docx 人话;空异常消息回落类型名;.pptx(标题/文字/表格/备注)
  ⑥ pdf_read_pages —— 总字符预算 + nextPage/next_pages 续读;加密 -> 明确报错;无文字层 -> OCR 提示
  ⑦ excel_read     —— 只读一次(不再非只读重开取数字格式);无 <dimension> 的表照读
  ⑧ browser        —— 模块 import json;get_text 选择器无匹配报错 + 截断;screenshot 走 encode_with_budget

Run with UTF-8:  python -X utf8 tests/smoke_office_read_v2.py
"""

import asyncio
import builtins
import io
import json
import os
import shutil
import sys
import tempfile
import time
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(_ROOT, "src"))

_TMP = tempfile.mkdtemp(prefix="acc_smoke_office_read_v2_")
_DATA = os.path.join(_TMP, "data")
os.makedirs(_DATA, exist_ok=True)
os.environ["WCW_DATA_DIR"] = _DATA

import ai_computer_control.server as server  # noqa: E402
import ai_computer_control.tools.filesystem as fs  # noqa: E402
import ai_computer_control.tools.memory as mem  # noqa: E402
import ai_computer_control.tools.office_read as oread  # noqa: E402
import ai_computer_control.tools.browser as browser  # noqa: E402

_FNS = {t.name: t.fn for t in server.mcp._tool_manager.list_tools()}
_FAILURES: list[str] = []


def check(cond: bool, msg: str):
    print(f"  [{'ok  ' if cond else 'FAIL'}] {msg}")
    if not cond:
        _FAILURES.append(msg)


def P(*parts):
    return os.path.join(_TMP, *parts)


def wb_bytes(path, data):
    with open(path, "wb") as f:
        f.write(data)


def section_edit_file():
    print("== ① edit_file: CRLF / BOM 保真 ==")
    bat = P("a.bat")
    wb_bytes(bat, b"@echo off\r\nset A=1\r\nset B=2\r\n")
    r = _FNS["edit_file"](path=bat, old_string="set A=1", new_string="set A=9")
    check(r.get("success") is True, f".bat 单行替换成功 (got {r})")
    check(open(bat, "rb").read() == b"@echo off\r\nset A=9\r\nset B=2\r\n", "单行替换后仍是 CRLF (回归: 曾整文件变 LF)")

    r = _FNS["edit_file"](path=bat, old_string="set A=9\nset B=2", new_string="set A=1\nset B=2\nset C=3")
    check(r.get("success") is True and open(bat, "rb").read() == b"@echo off\r\nset A=1\r\nset B=2\r\nset C=3\r\n",
          "多行 old/new 用 LF 书写也能匹配 CRLF 文件,新增行同样是 CRLF")

    ps1 = P("b.ps1")
    wb_bytes(ps1, b"\xef\xbb\xbf" + "Write-Host '你好'\r\nWrite-Host 2\r\n".encode("utf-8"))
    r = _FNS["edit_file"](path=ps1, old_string="Write-Host 2", new_string="Write-Host 3")
    got = open(ps1, "rb").read()
    check(got == b"\xef\xbb\xbf" + "Write-Host '你好'\r\nWrite-Host 3\r\n".encode("utf-8"),
          f".ps1 BOM 只有一个且 CRLF 保留 (got {got[:12]!r}...)")
    check(got.count(b"\xef\xbb\xbf") == 1, "BOM 不双写 (回归: 曾变成 EF BB BF EF BB BF)")

    nobom = P("c.txt")
    wb_bytes(nobom, "无 BOM 文件\n第二行\n".encode("utf-8"))
    _FNS["edit_file"](path=nobom, old_string="第二行", new_string="第三行")
    check(open(nobom, "rb").read() == "无 BOM 文件\n第三行\n".encode("utf-8"), "无 BOM 的 LF 文件不被加 BOM、不被改行尾")

    sig = P("d.txt")
    wb_bytes(sig, b"\xef\xbb\xbf" + b"a\nb\n")
    _FNS["edit_file"](path=sig, old_string="b", new_string="c", encoding="utf-8-sig")
    check(open(sig, "rb").read() == b"\xef\xbb\xbfa\nc\n", "encoding=utf-8-sig 也只有一个 BOM")

    miss = _FNS["edit_file"](path=bat, old_string="没有这段", new_string="x")
    check("未出现" in (miss.get("error") or ""), "找不到时仍给诊断错误")

    # 混用换行:一个散落的 LF 不该让所有多行 LF old_string 失效;未改动的行保持原换行
    mx = P("mixed.txt")
    wb_bytes(mx, b"a\r\nb\r\nc\r\nd\ne\r\nf\r\n")
    r = _FNS["edit_file"](path=mx, old_string="b\nc", new_string="B\nC\nC2")
    check(r.get("success") is True and open(mx, "rb").read() == b"a\r\nB\r\nC\r\nC2\r\nd\ne\r\nf\r\n",
          f"混用换行文件: CRLF 处的多行 LF 编辑命中,new 取该处的 CRLF (got {open(mx, 'rb').read()!r})")
    r = _FNS["edit_file"](path=mx, old_string="d\ne", new_string="D\nE")
    check(r.get("success") is True and open(mx, "rb").read() == b"a\r\nB\r\nC\r\nC2\r\nD\nE\r\nf\r\n",
          "混用换行文件: LF 处的多行编辑也命中,保持 LF")
    wb_bytes(mx, b"a\r\nb\r\nc\r\nd\ne\r\nf\r\n")
    r = _FNS["edit_file"](path=mx, old_string="c\nd\ne", new_string="x")
    err = r.get("error") or ""
    check("TypeError" not in err and "CRLF 与 LF 混用" in err,
          f"跨越混用换行处的片段: 给出换行指引而不是裸 TypeError (got {err[:80]!r})")
    # 诊断路径的哨兵字符不再触发 TypeError
    wb_bytes(mx, b"hello wor\nnext")
    r = _FNS["edit_file"](path=mx, old_string="hello world", new_string="x")
    check("未出现" in (r.get("error") or "") and "TypeError" not in (r.get("error") or ""),
          f"old_string 越过文件尾的诊断不再抛 TypeError (got {(r.get('error') or '')[:80]!r})")
    wb_bytes(mx, b"alpha\nbeta\n")
    r = _FNS["edit_file"](path=mx, old_string="alpha\nbeta\ngamma", new_string="x")
    check("未出现" in (r.get("error") or "") and "TypeError" not in (r.get("error") or ""),
          "多行 old_string 越过文件尾的诊断不再抛 TypeError")


def section_read_file():
    print("\n== ② read_file: BOM / UTF-16 / GBK / 二进制 / 噪声 ==")
    f = P("bom.txt")
    wb_bytes(f, b"\xef\xbb\xbf" + "你好 BOM".encode("utf-8"))
    r = _FNS["read_file"](path=f)
    check(r.get("content") == "你好 BOM" and not r["content"].startswith("﻿"), f"UTF-8 BOM 被剥掉 (got {r.get('content')!r})")

    f = P("u16.txt")
    wb_bytes(f, "hello 世界\r\n".encode("utf-16"))  # 带 BOM
    r = _FNS["read_file"](path=f)
    check(r.get("content") == "hello 世界\r\n" and "\x00" not in r.get("content", ""),
          f"UTF-16(BOM) 正确解码 (got {r.get('content')!r}, enc={r.get('encoding_used')})")
    f = P("u16nb.txt")
    wb_bytes(f, "line one\nline two\n".encode("utf-16-le"))  # 无 BOM
    r = _FNS["read_file"](path=f)
    check(r.get("content") == "line one\nline two\n", f"无 BOM 的 UTF-16-LE 也识别 (got {r.get('content')!r})")

    f = P("gbk.txt")
    wb_bytes(f, "中文内容,来自记事本 ANSI".encode("gbk"))
    r = _FNS["read_file"](path=f)
    check("中文内容" in r.get("content", "") and "�" not in r.get("content", ""),
          f"GBK 文件回退解码无乱码 (enc={r.get('encoding_used')})")
    check(bool(r.get("encoding_fallback")), "标注了 encoding_fallback")

    f = P("bin.bin")
    wb_bytes(f, bytes([0x89, 0x50, 0x4E, 0x47, 0, 0, 0, 13]) + os.urandom(2000).replace(b"\x00", b"\x01") + b"\x00\x00")
    r = _FNS["read_file"](path=f)
    check(r.get("ok") is False and r.get("binary") is True and r.get("hint"), f"二进制拒读并给提示 (got keys {sorted(r)})")

    f = P("cn.txt")
    wb_bytes(f, ("这是一份普通的中文说明文档,里面有很多汉字和标点符号。" * 30).encode("utf-8"))
    r = _FNS["read_file"](path=f)
    check("non_ascii" not in r, f"普通中文不附 non_ascii 噪声 (got {len(json.dumps(r.get('non_ascii', {}), ensure_ascii=False))} 字节)")

    f = P("sus.txt")
    wb_bytes(f, "x = 1 + 2\nprint(（\"hi\"）)\n价格 — 10 元\n".encode("utf-8"))
    r = _FNS["read_file"](path=f)
    chars = {s["char"] for s in r.get("non_ascii", {}).get("samples", [])}
    check({" ", "—"} <= chars, f"易混字符(nbsp / em dash)仍被报告 (got {chars})")
    check("（" in chars, "紧贴 ASCII 的全角括号仍被报告")


def section_read_file_truncation():
    print("\n== ②b read_file: max_bytes 落在 UTF-8 字符中间 ==")
    f = P("cjk_long.txt")
    text = "你好世界,这是一份很长的中文文档。" * 400
    wb_bytes(f, text.encode("utf-8"))
    bad = []
    for mb in range(1000, 1007):
        r = _FNS["read_file"](path=f, max_bytes=mb)
        if r.get("encoding_used") != "utf-8" or "encoding_fallback" in r or not text.startswith(r.get("content", "x")):
            bad.append((mb, r.get("encoding_used")))
    check(not bad, f"截断点落在多字节字符中间也仍按 UTF-8 解码(不整文件退成 GBK) (bad: {bad})")
    r = _FNS["read_file"](path=f, max_bytes=1001)
    check(r.get("truncated") is True and text.startswith(r["content"]) and len(r["content"].encode("utf-8")) >= 998
          and "�" not in r["content"], "截断后内容是完整字符的前缀(无 U+FFFD)")
    g = P("gbk_long.txt")
    wb_bytes(g, ("中文内容来自记事本。" * 300).encode("gbk"))
    r = _FNS["read_file"](path=g, max_bytes=1001)
    check(r.get("encoding_fallback") and "中文内容" in r.get("content", ""), "真 GBK 文件截断时仍正常回退 GBK")
    check(fs._trim_partial_utf8("你".encode("utf-8")[:2]) == b"" and fs._trim_partial_utf8(b"ab") == b"ab"
          and fs._trim_partial_utf8("你好".encode("utf-8")) == "你好".encode("utf-8"), "_trim_partial_utf8 只丢残缺尾部")


def section_list_directory():
    print("\n== ③ list_directory ==")
    root = P("tree")
    os.makedirs(os.path.join(root, ".git", "objects"))
    os.makedirs(os.path.join(root, "src"))
    for n in (".git/config", ".git/objects/ab", "src/main.py", "README.md", ".env"):
        wb_bytes(os.path.join(root, *n.split("/")), b"x")
    made_link = True
    try:
        os.symlink(os.path.join(root, "nonexistent-target"), os.path.join(root, "src", "broken"))
    except (OSError, NotImplementedError):
        made_link = False
    r = _FNS["list_directory"](path=root, recursive=True)
    names = {e["name"].replace("\\", "/") for e in r.get("entries", [])}
    check(r.get("error") is None, f"递归列目录不因坏链接整体失败 (got {r.get('error')})")
    check(not any(n.startswith(".git") for n in names) and ".env" not in names,
          f"include_hidden=False 不列 .git/** 与隐藏文件 (got {sorted(names)})")
    check({"src", "src/main.py", "README.md"} <= names, "普通条目都在")
    if made_link:
        check(r.get("skipped") == 1, f"坏符号链接被跳过并计数 skipped=1 (got {r.get('skipped')})")
    else:
        print("  [skip] 本环境不能建符号链接(Windows 无权限),跳过坏链接断言")
    r2 = _FNS["list_directory"](path=root, recursive=True, include_hidden=True)
    names2 = {e["name"].replace("\\", "/") for e in r2.get("entries", [])}
    check(".git/config" in names2 and ".env" in names2, "include_hidden=True 时 .git 内容可见")

    keys = list(r)
    check(keys.index("total") < keys.index("entries") and [k for k in keys if k != "ok"][-1] == "entries",
          f"total 等摘要键排在 entries 之前(宿主按字符硬切时不被截没) keys={keys}")

    big = P("flat")
    os.makedirs(big)
    for i in range(1500):
        open(os.path.join(big, f"f{i:04d}.txt"), "w").close()
    r = _FNS["list_directory"](path=big)
    check(r.get("capped") is True and len(r["entries"]) == 1000 and r["total"] == 1000,
          f"非递归列目录也封顶 1000 (got total={r.get('total')}, capped={r.get('capped')})")
    keys = list(r)
    check(keys.index("capped") < keys.index("entries"), "capped 在 entries 之前")
    r = _FNS["list_directory"](path=big, limit=10)
    check(len(r["entries"]) == 10 and r.get("capped") is True, "limit 参数生效")
    r = _FNS["list_directory"](path=os.path.join(root, "src"))
    check(r.get("capped") is None, "没封顶不带 capped 键")


def section_memory():
    print("\n== ④ memory: 瞬时 OSError 不隔离健康库 ==")
    store = os.path.join(_DATA, "memory.json")
    r = _FNS["memory_save"](key="pref", content="深色主题")
    check(r.get("success") is True and os.path.exists(store), "先存一条")

    real_open = builtins.open
    state = {"left": 1}

    def flaky_open(file, mode="r", *a, **k):
        if os.path.basename(str(file)) == "memory.json" and "r" in mode and "b" not in mode and state["left"] > 0:
            state["left"] -= 1
            raise PermissionError(13, "The process cannot access the file because it is being used by another process")
        return real_open(file, mode, *a, **k)

    builtins.open = flaky_open
    try:
        r = _FNS["memory_read"](key="pref")   # 第一次读被锁,短重试后应读到
    finally:
        builtins.open = real_open
    check(r.get("found") is True and r.get("content") == "深色主题", f"瞬时锁后重试读到 (got {r})")
    check(os.path.exists(store) and not os.path.exists(store + ".corrupt"), "健康库没被改名成 .corrupt")

    state["left"] = 10_000   # 持续被锁
    builtins.open = flaky_open
    try:
        r = _FNS["memory_read"](key="pref")
        rs = _FNS["memory_save"](key="pref2", content="x")
    finally:
        builtins.open = real_open
    check(r.get("error") and "found" not in r, f"持续被锁: 返回明确错误而不是空库 (got {r})")
    check(rs.get("error"), "持续被锁时 memory_save 不会用空库覆盖")
    check(os.path.exists(store) and not os.path.exists(store + ".corrupt"), "持续被锁也不隔离")
    check(_FNS["memory_read"](key="pref").get("found") is True, "锁释放后数据完好")

    with open(store, "w", encoding="utf-8") as f:
        f.write("{not json")
    r = _FNS["memory_read"](key="pref")
    check(r.get("found") is False and os.path.exists(store + ".corrupt"), "真解析失败才隔离成 .corrupt")

    # _save: per-process tmp name, transient sharing violation retried, no orphan tmp on failure
    os.remove(store + ".corrupt") if os.path.exists(store + ".corrupt") else None
    os.remove(store) if os.path.exists(store) else None
    names = []
    real_replace = os.replace
    flaky = {"left": 2}

    def flaky_replace(src, dst):
        names.append(os.path.basename(src))
        if flaky["left"] > 0:
            flaky["left"] -= 1
            e = PermissionError(13, "sharing violation")
            e.winerror = 32
            raise e
        return real_replace(src, dst)
    os.replace = flaky_replace
    try:
        r = _FNS["memory_save"](key="retry", content="x")
    finally:
        os.replace = real_replace
    check(r.get("success") is True and len(names) == 3, f"memory_save 遇到瞬时占用(WinError 32)重试后成功 (replace 调用 {len(names)} 次, got {r})")
    check(all(str(os.getpid()) in n for n in names), f"memory 临时文件名含进程号(两个 ACC 进程不共用一个 .tmp) (got {names[:1]})")
    check(not [n for n in os.listdir(_DATA) if n.endswith(".tmp")], "成功后没有残留 .tmp")


try:
    import reportlab  # noqa: F401
    HAVE_RL = True
except Exception:
    HAVE_RL = False   # pdf-export 是可选 extra;CI 的最小安装可能没有 -> 相关断言跳过


def _make_pdf(path, pages=1, text_lines=0, encrypt=None, blank=False):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path, encrypt=encrypt) if encrypt else canvas.Canvas(path)
    for p in range(1, pages + 1):
        if not blank:
            c.setFont("Helvetica", 9)
            c.drawString(50, 800, f"PAGE-{p}-HEADER")
            for i in range(text_lines):
                c.drawString(50, 780 - i * 10, f"page {p} line {i} lorem ipsum dolor sit amet consectetur")
        c.showPage()
    c.save()


def section_read_document():
    print("\n== ⑤ read_document ==")
    from docx import Document

    d = Document()
    d.add_heading("第一章 概述", level=1)
    d.add_paragraph("这是正文第一段。")
    d.add_heading("1.1 背景", level=2)
    d.add_paragraph("要点甲", style="List Bullet")
    d.add_paragraph("要点乙", style="List Bullet")
    d.add_paragraph("步骤一", style="List Number")
    d.add_paragraph("步骤二", style="List Number")
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2))
    t.cell(0, 0).text = "合并表头"
    for c, v in enumerate(("名称", "数量", "备注")):
        t.cell(1, c).text = v
    for c, v in enumerate(("苹果", "3", "含|竖线")):
        t.cell(2, c).text = v
    d.add_paragraph("表格之后的结尾段。")
    docx_path = P("struct.docx")
    d.save(docx_path)
    r = _FNS["read_document"](path=docx_path)
    c = r.get("content", "")
    check(r.get("ok") is True and "# 第一章 概述" in c and "## 1.1 背景" in c, f"标题保留为 markdown # (got {c[:80]!r})")
    check("- 要点甲" in c and "1. 步骤一" in c and "2. 步骤二" in c, "列表保留(项目符号 / 编号)")
    check(c.count("合并表头") == 1, f"合并单元格文字不重复 (count={c.count('合并表头')})")
    check("| 名称 | 数量 | 备注 |" in c and "| --- |" in c and "含\\|竖线" in c, "表格为 markdown 行,竖线被转义")
    check(c.index("- 要点乙") < c.index("合并表头") < c.index("表格之后的结尾段"), "表格在文档中的原位置(不是堆到末尾)")
    check([h["text"] for h in r.get("headings", [])] == ["第一章 概述", "1.1 背景"], "headings 大纲")
    check("non_ascii" not in r, "普通中文 docx 不附 non_ascii 噪声")

    d = Document()
    for i in range(400):
        d.add_paragraph(f"第{i:04d}段:" + "内容" * 40)
    big = P("big.docx")
    d.save(big)
    r1 = _FNS["read_document"](path=big, max_chars=5000)
    check(r1.get("truncated") is True and r1.get("next_offset") == 5000 and len(r1["content"]) == 5000
          and r1.get("total_chars", 0) > 30000, f"docx 有字符上限并给 next_offset (total={r1.get('total_chars')})")
    got, off = "", 0
    for _ in range(100):
        rr = _FNS["read_document"](path=big, max_chars=20000, offset=off)
        got += rr["content"]
        off = rr.get("next_offset")
        if off is None:
            break
    full = _FNS["read_document"](path=big, max_chars=200000)["content"]
    check(got == full and "第0399段" in got, "按 next_offset 续读拼回全文")

    # --- 加密 / 扫描 / 损坏 ---
    if HAVE_RL:
        enc = P("enc.pdf")
        _make_pdf(enc, pages=1, text_lines=3, encrypt="secret")
        r = _FNS["read_document"](path=enc)
        check(r.get("ok") is False and "加密" in (r.get("error") or ""),
              f"加密 PDF -> ok:false + 明确原因 (回归: 曾 ok:true error:'') got={r}")
        scan = P("scan.pdf")
        _make_pdf(scan, pages=2, blank=True)
        r = _FNS["read_document"](path=scan)
        check(r.get("ok") is True and "OCR" in (r.get("note") or ""),
              f"扫描件 PDF: ok 且 note 建议 OCR (got note={r.get('note')!r})")
    else:
        print("  [skip] reportlab 未装,跳过 PDF 夹具断言")

    fake = P("locked.xlsx")
    wb_bytes(fake, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
    r = _FNS["read_document"](path=fake)
    check(r.get("ok") is False and "密码" in (r.get("error") or "") and "zip" not in (r.get("error") or "").lower(),
          f"加密 xlsx(OLE 头) -> 提示密码而非 'not a zip file' (got {r.get('error')!r})")
    bad = P("bad.docx")
    wb_bytes(bad, b"this is not a zip at all")
    r = _FNS["read_document"](path=bad)
    check(r.get("ok") is False and "不是有效的 .docx" in (r.get("error") or ""), f"损坏 docx 人话 (got {r.get('error')!r})")
    check(oread.describe_read_error(Exception(), "x.docx") == "Exception", "空异常消息回落到类型名")
    r = _FNS["read_document"](path=P("x.txt"))
    check(r.get("ok") is False and "read_file" in (r.get("error") or ""), "不支持的格式给出去向提示")

    # --- pptx ---
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception:
        print("  [skip] python-pptx 未装,跳过 .pptx")
    else:
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[1])
        s1.shapes.title.text = "季度汇报"
        s1.placeholders[1].text = "第一条要点\n第二条要点"
        s1.notes_slide.notes_text_frame.text = "讲者备注:强调增长"
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        s2.shapes.title.text = "数据表"
        tb = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for (r_, c_), v in {(0, 0): "指标", (0, 1): "值", (1, 0): "营收", (1, 1): "12%"}.items():
            tb.cell(r_, c_).text = v
        pp = P("deck.pptx")
        prs.save(pp)
        r = _FNS["read_document"](path=pp)
        c = r.get("content", "")
        check(r.get("ok") is True and r.get("slides") == 2 and "## Slide 1: 季度汇报" in c and "第二条要点" in c
              and "[备注] 讲者备注:强调增长" in c and "## Slide 2: 数据表" in c and "| 营收 | 12% |" in c,
              f".pptx: 标题/文字/备注/表格 (got {c!r})")
        # 一个 python-pptx 无法归类 shape_type 的形状(无 prstGeom、非 txBox)不该让整套 deck 读失败
        s3 = prs.slides.add_slide(prs.slide_layouts[6])
        tb3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        tb3.text_frame.text = "无法归类的形状里的文字"
        sp = tb3._element
        if "txBox" in sp.nvSpPr.cNvSpPr.attrib:
            del sp.nvSpPr.cNvSpPr.attrib["txBox"]
        for g in sp.spPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom"):
            sp.spPr.remove(g)
        pp2 = P("deck_odd.pptx")
        prs.save(pp2)
        r = _FNS["read_document"](path=pp2)
        c = r.get("content", "")
        check(r.get("ok") is True and r.get("slides") == 3 and "## Slide 1: 季度汇报" in c and "无法归类的形状里的文字" in c,
              f".pptx: shape_type 抛错的形状不使整份读取失败,其余内容照读 (got {r.get('error') or c[-60:]!r})")


def section_pdf_pages():
    print("\n== ⑥ pdf_read_pages ==")
    if not HAVE_RL:
        print("  [skip] reportlab 未装,跳过")
        return
    pdf = P("long.pdf")
    _make_pdf(pdf, pages=30, text_lines=60)
    r = _FNS["pdf_read_pages"](path=pdf, pages="1-30", max_chars=12000)
    total = sum(p["chars"] for p in r.get("pages", []))
    check(r.get("ok") is True and total <= 12000 and r.get("truncated") is True,
          f"总字符预算生效 (chars={total}, pages={len(r.get('pages', []))}, truncated={r.get('truncated')})")
    check(r.get("nextPage") == len(r["pages"]) + 1 and r.get("next_pages") == f"{r['nextPage']}-30",
          f"给出 nextPage / next_pages (got {r.get('nextPage')}, {r.get('next_pages')})")
    seen = [p["page"] for p in r["pages"]]
    spec, guard = r.get("next_pages"), 0
    while spec and guard < 40:
        rr = _FNS["pdf_read_pages"](path=pdf, pages=spec, max_chars=12000)
        seen += [p["page"] for p in rr["pages"]]
        spec = rr.get("next_pages")
        guard += 1
    check(seen == list(range(1, 31)), f"按 next_pages 续读恰好覆盖 1..30 各一次 (got {seen})")
    r = _FNS["pdf_read_pages"](path=pdf, pages="1-3")
    check("truncated" not in r and "next_pages" not in r, "预算够用时不带续读键")
    t0 = time.time()
    try:
        oread._parse_pages("1-999999999", 30)
        ok_raise = False
    except ValueError as ve:
        ok_raise = "越界" in str(ve)
    check(ok_raise and time.time() - t0 < 0.5, "页码范围先做边界检查再展开('1-999999999' 立刻报越界,不建巨大列表)")
    check(oread._parse_pages("3-1,2,7-8", 30) == [1, 2, 3, 7, 8] and oread._parse_pages("5", 30) == [5],
          "合法页码范围解析不变")

    enc = P("enc2.pdf")
    _make_pdf(enc, pages=1, text_lines=3, encrypt="secret")
    r = _FNS["pdf_read_pages"](path=enc)
    check(r.get("ok") is False and "加密" in (r.get("error") or ""), f"加密 PDF 明确报错 (got {r.get('error')!r})")

    scan = P("scan2.pdf")
    _make_pdf(scan, pages=2, blank=True)
    r = _FNS["pdf_read_pages"](path=scan, pages="1-2")
    check(r.get("ok") is True and r.get("no_text_pages") == [1, 2] and "ocr_image" in (r.get("note") or ""),
          f"无文字层页给 no_text_pages + OCR 提示 (got {r.get('note')!r})")


def section_excel():
    print("\n== ⑦ excel_read ==")
    import datetime
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["名称", "金额", "日期"])
    for i in range(300):
        ws.append([f"n{i}", i * 10.5, datetime.date(2024, 1, 1) + datetime.timedelta(days=i)])
    ws["B2"].number_format = "¥#,##0.00"
    ws["C2"].number_format = "yyyy-mm-dd"
    xl = P("fmt.xlsx")
    wb.save(xl)

    calls = []
    real = oread._load_workbook_bounded

    def spy(path, timeout=20, **kw):
        calls.append(dict(kw))
        return real(path, timeout=timeout, **kw)

    oread._load_workbook_bounded = spy
    try:
        r = _FNS["excel_read"](path=xl, max_rows=50)
    finally:
        oread._load_workbook_bounded = real
    check(r.get("ok") is True and r.get("number_formats") == {"B": "¥#,##0.00", "C": "yyyy-mm-dd"},
          f"数字格式读到 (got {r.get('number_formats')})")
    check(len(calls) == 1 and calls[0].get("read_only") is True,
          f"只开一次工作簿且是 read_only (回归: 曾再整本非只读重开) calls={calls}")
    check(r.get("truncated", {}).get("rows_total") == 301 and r["truncated"].get("next_range") == "A51:C301",
          f"截断信息带 next_range (got {r.get('truncated')})")

    # 无 <dimension> 的表
    nodim = P("nodim.xlsx")
    zin = zipfile.ZipFile(xl)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        import re
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            zo.writestr(it, data)
    wb_bytes(nodim, out.getvalue())
    r = _FNS["excel_read"](path=nodim, max_rows=20)
    check(r.get("ok") is True and len(r.get("data", [])) == 19 and r.get("headers", [None])[0] == "名称",
          f"无 <dimension> 的表照读 (回归: 曾读成空表) rows={len(r.get('data', []))}")
    check(r["sheets"][0]["rows"] is None, "概要里未知维度报 None 而不是 0")
    check(r.get("truncated", {}).get("rows_total") is None and r["truncated"].get("next_start_row") == 21,
          f"未知总行数如实标注 (got {r.get('truncated')})")
    r = _FNS["excel_read"](path=nodim, max_rows=1000)
    check(len(r.get("data", [])) == 300 and r.get("truncated") is None, "读完整张表且不误报截断")

    r = _FNS["excel_read"](path=P("locked.xlsx"))
    check(r.get("ok") is False and "密码" in (r.get("error") or ""), f"加密 xlsx 提示密码 (got {r.get('error')!r})")


class _FakeEl:
    def __init__(self, text):
        self._t = text

    async def inner_text(self):
        return self._t


class _FakePage:
    url = "http://example.test/"

    def __init__(self, body="", sel=None, png=None):
        self.body, self.sel, self.png = body, sel or {}, png

    async def query_selector(self, s):
        return _FakeEl(self.sel[s]) if s in self.sel else None

    async def inner_text(self, s):
        return self.body

    async def screenshot(self, type="png"):
        return self.png

    async def title(self):
        return "T"


def section_fetch_html():
    print("\n== ⑨ fetch: HTML -> text 线性 / 标题 / Location 大小写 ==")
    import ai_computer_control.tools.web_fetch as wf
    t0 = time.time()
    text, _title = wf._html_to_text("<pre>a" + " " * 200000 + "b</pre>", "http://example.test/")
    dt = time.time() - t0
    check(dt < 2.0 and text.startswith("a") and text.endswith("b"), f"200KB 空格串的 <pre> 线性处理 (耗时 {dt:.2f}s)")
    t0 = time.time()
    wf._html_to_text("<pre>a" + " \t" * 100000 + "b</pre>", "http://example.test/")
    check(time.time() - t0 < 2.0, "空格+制表符混排同样线性")
    text, title = wf._html_to_text(
        "<html><head><title>My Site</title></head><body><svg><title>menu icon</title></svg>"
        "<p>hello</p><svg><title>close</title></svg></body></html>", "http://example.test/")
    check(title == "My Site" and text == "hello", f"内联 <svg><title> 不污染页面标题/正文 (got {title!r}, {text!r})")
    _t, title = wf._html_to_text("<title>A</title><title>B</title><p>x</p>", "http://example.test/")
    check(title == "A", "只取第一个 <title>")

    calls = []

    def fake_once(url, tmo, budget):
        calls.append(url)
        if len(calls) == 1:
            return 302, {"LOCATION": "http://example.test/next"}, b"", None
        return 200, {"content-type": "text/plain"}, b"done", None
    old = (wf._check_url, wf._fetch_once)
    wf._check_url, wf._fetch_once = (lambda u: None), fake_once
    try:
        r = _FNS["fetch"](url="http://example.test/")
    finally:
        wf._check_url, wf._fetch_once = old
    check(r.get("ok") is True and calls[-1] == "http://example.test/next", f"大写 LOCATION 头也跟随重定向 (got {r.get('error') or r.get('url')})")


def section_browser():
    print("\n== ⑧ browser 纯逻辑(playwright 未装, 用假 page) ==")
    check(hasattr(browser, "json"), "browser 模块 import 了 json (回归: 曾 NameError 被吞)")
    page = _FakePage(body="长" * 50000, sel={"#a": "hello"})

    async def ensure():
        return page

    old_av, old_ens = browser._AVAILABLE, browser._ensure_browser
    browser._AVAILABLE, browser._ensure_browser = True, ensure
    try:
        run = asyncio.run
        r = run(browser.browser_get_text(selector="#missing"))
        check(r.get("error") and "matched no element" in r["error"], f"选择器无匹配 -> 明确错误 (got {r})")
        r = run(browser.browser_get_text(selector="#a"))
        check(r.get("text") == "hello" and r.get("truncated") is False, "命中时正常返回")
        r = run(browser.browser_get_text())
        check(len(r["text"]) == 20000 and r["truncated"] is True and r["total_chars"] == 50000,
              f"get_text 默认封顶 20000 并标 truncated (got {len(r['text'])})")
        r = run(browser.browser_get_text(max_chars=100))
        check(len(r["text"]) == 100, "max_chars 参数生效")

        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (2400, 1200), (10, 120, 200)).save(buf, format="PNG")
        page.png = buf.getvalue()
        r = run(browser.browser_screenshot())
        check(r.get("width") == 1280 and r.get("scale") == round(1280 / 2400, 4) and r.get("format") == "png",
              f"screenshot 默认按 max_width=1280 缩放 (got w={r.get('width')}, scale={r.get('scale')})")
        r = run(browser.browser_screenshot(max_width=0, format="jpeg", quality=50))
        check(r.get("width") == 2400 and r.get("format") == "jpeg" and r.get("url") == "http://example.test/",
              "max_width=0 保持原尺寸, format=jpeg 生效")
    finally:
        browser._AVAILABLE, browser._ensure_browser = old_av, old_ens


def main() -> int:
    try:
        for sec in (section_edit_file, section_read_file, section_list_directory, section_memory,
                    section_read_document, section_pdf_pages, section_excel, section_browser,
                    section_read_file_truncation, section_fetch_html):
            try:
                sec()
            except Exception as e:  # 一节崩了不遮住其它节的结果
                check(False, f"{sec.__name__} 抛异常: {type(e).__name__}: {e}")
    finally:
        shutil.rmtree(_TMP, ignore_errors=True)
    print()
    if _FAILURES:
        print(f"OFFICE READ V2 SMOKE: FAIL ({len(_FAILURES)})")
        for m in _FAILURES:
            print("  -", m)
        return 1
    print("OFFICE READ V2 SMOKE: ALL PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
