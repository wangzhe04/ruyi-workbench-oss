"""Round-trip tests for lossless pagination, spreadsheet types, and Word pagination rules."""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
os.environ['WCW_DATA_DIR'] = tempfile.mkdtemp(prefix='acc-office-quality-')
import ai_computer_control.server as server
from openpyxl import load_workbook
from pptx import Presentation
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

tools = {tool.name: tool.fn for tool in server.mcp._tool_manager.list_tools()}
with tempfile.TemporaryDirectory(prefix='office-quality-') as tmp:
    root = Path(tmp)
    portrait = root / 'portrait.png'
    Image.new('RGB', (200, 800), 'blue').save(portrait)
    ppt = root / 'slides.pptx'
    bullets = [f'保留要点 {i}' for i in range(23)]
    rows = [[f'记录 {i}', i] for i in range(25)]
    result = tools['write_pptx'](str(ppt), [
        {'type': 'content', 'title': '完整内容', 'bullets': bullets},
        {'type': 'table', 'title': '完整数据', 'headers': ['项目', '金额'], 'rows': rows},
        {'type': 'image', 'title': '等比居中', 'image_path': str(portrait)},
    ])
    assert result.get('success'), result
    assert result['paginated'] and result['slides'] == 9, result
    prs = Presentation(ppt)
    all_text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    for bullet in bullets:
        assert bullet in all_text, bullet
    tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert sum(len(table.rows) - 1 for table in tables) == 25
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= prs.slide_height + 10, 'shape below slide'
    pic = next(shape for shape in prs.slides[-1].shapes if shape.shape_type == 13)
    assert abs(2 * pic.left + pic.width - prs.slide_width) < 10

    book = root / 'data.xlsx'
    result = tools['write_excel'](str(book), [['00123', '123456789012345678', '12%', '150%', '=1+2'],
                                             ['00001', '000000000000000001', '-0.5%', '0%', '=SUM(1,2)']],
                                  headers=['编号', '标识符', '比例', '增长率', '计算'])
    assert result.get('success'), result
    assert result['formula_status'] == 'not_calculated'
    assert tools['excel_beautify'](str(book)).get('success')
    wb = load_workbook(book)
    ws = wb.active
    assert ws['A2'].value == '00123' and ws['A2'].data_type == 's'
    assert ws['B2'].value == '123456789012345678'
    assert ws['C2'].value == .12 and ws['C2'].number_format == '0.0%'
    assert ws['D2'].value == 1.5 and ws['D2'].number_format == '0.0%'
    assert ws['C3'].value == -.005 and ws['D3'].value == 0
    assert ws['E2'].value == '=1+2' and ws['E2'].data_type == 'f'
    assert ws.freeze_panes == 'A2' and ws.auto_filter.ref == 'A1:E3'
    assert ws.print_title_rows == '$1:$1'
    wb.close()

    doc = root / 'report.docx'
    result = tools['write_document'](str(doc), '# 标题\n4. 第四项\n10. 第十项\nTABLE: 项目 | 金额\n| A | 100\n', page_numbers=True)
    assert result.get('success'), result
    document = Document(doc)
    numbered = [p for p in document.paragraphs if p.style.name == 'List Number']
    assert [p.text for p in numbered] == ['第四项', '第十项']
    assert document.styles['Heading 1'].paragraph_format.keep_with_next
    assert document.styles['Normal'].paragraph_format.widow_control
    table = document.tables[0]
    assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    assert all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows)
print('OFFICE QUALITY ROUND-TRIP: ALL PASS')
